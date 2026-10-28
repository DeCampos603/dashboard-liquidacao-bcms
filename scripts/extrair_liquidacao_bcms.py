# -*- coding: utf-8 -*-
"""
Extrator de Liquidação & Restos a Pagar do BCMS (SIAFI / Tesouro Gerencial)
Batalhão Central de Manutenção e Suprimento (BCMS) · Base de Apoio Logístico (Ba Ap Log)

REGRA: nenhum valor é estimado. O que a fonte não traz fica explicitamente como
"sem dado" (None) e o painel exibe essa limitação.

Fontes reais:
  - EMPENHO (Google Drive, link publicado): NE, valor da NE, itens, contas 622110000/622130100.
    NÃO traz liquidação por NE (sem as contas 622920101/622920103).
  - RPNP (cache data/rpnp_cache.xlsx): contas 531110100, 631100000..631400000 por ação.
  - Totais de liquidação por UG: Dashboard-Credito-BCMS (om_totais_siafi), consolidado SIAFI.
"""

import sys
import json
import re
import datetime
import pathlib
import urllib.request
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

SCRIPT_DIR = pathlib.Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
DATA_DIR = PROJECT_DIR / "data"
CACHE_XLSX = DATA_DIR / "empenhos_cache.xlsx"
RPNP_XLSX = DATA_DIR / "rpnp_cache.xlsx"
TOTAIS_JSON = DATA_DIR / "totais_siafi_ug.json"
OUTPUT_JSON = DATA_DIR / "liquidacao_data.json"

LOCAL_EMPENHO = PROJECT_DIR.parent / "Acervo-Empenhos-BCMS-2026" / "EMPENHO.xlsx"
DEFAULT_EMPENHOS_URL = "https://docs.google.com/spreadsheets/d/e/2PACX-1vSTGXyEQ2v4UYga0nkP7ypbyumOZd68m32UZwqVVQhx-bHBACkBSj7v4RiNWZydIQ/pub?output=xlsx"
UGS = ("160329", "167329")
UG_NOME = {"160329": "BCMS (OGU)", "167329": "BCMS (FEx)"}
RPNP_DATA_REF = "02/10/2026"


def num(x):
    try:
        return float(x)
    except Exception:
        return 0.0


def format_cnpj(c):
    d = re.sub(r"\D", "", str(c or ""))
    if len(d) == 14:
        return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"
    if len(d) == 11:
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    return str(c or "-")


def obter_planilha_empenhos():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    print("Verificando atualização no Google Drive...")
    req = urllib.request.Request(DEFAULT_EMPENHOS_URL, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            content = resp.read()
        if len(content) > 5000 and content[:2] == b"PK":
            CACHE_XLSX.write_bytes(content)
            print(f"-> Planilha baixada do Drive ({len(content) / 1024:.1f} KB).")
            return CACHE_XLSX, "Google Drive (baixada agora)"
    except Exception as e:
        print(f"[AVISO] Falha no download ({e}).")
    if CACHE_XLSX.exists() and CACHE_XLSX.stat().st_size > 5000:
        print("-> Usando cache versionado.")
        return CACHE_XLSX, "Cache versionado (Drive indisponível)"
    if LOCAL_EMPENHO.exists():
        return LOCAL_EMPENHO, "Arquivo local do acervo"
    raise FileNotFoundError("Planilha de empenhos indisponível.")


def extrair_empenhos(xlsx_path, corte):
    """Uma NE = linha com col 4 preenchida; as linhas seguintes (col 4 vazia) são itens dela."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["EMPENHO"] if "EMPENHO" in wb.sheetnames else wb.active
    nes, cur = {}, None
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[4]:
            raw = str(r[4])
            cur = None
            if not raw.startswith(UGS):
                continue
            ug, ne_num = raw[:6], raw[11:]
            key = f"{ug}_{ne_num}"
            emissao = str(r[2]).strip() if r[2] else ""
            aging = 0
            try:
                p = emissao.split("/")
                aging = (corte - datetime.date(int(p[2]), int(p[1]), int(p[0]))).days
            except Exception:
                pass
            desc = str(r[20]).strip() if r[20] else ""
            m_proc = re.search(r"(PE\s*[\d/_-]+|PREG[ÃA]O\s*[\d/_-]+|DISPENSA\s*[\d/_-]+|DIEX\s*[\d/_-]+)", desc, re.I)
            m_nc = re.search(r"(20\d{2}NC\d{6})", desc)
            valor = round(num(r[7]), 2)
            nes[key] = {
                "key": key, "ne_num": ne_num, "ne_full": raw, "ug": ug, "ug_nome": UG_NOME[ug],
                "emissao": emissao, "idade_dias": max(0, aging),
                "cnpj": format_cnpj(r[5]),
                "fornecedor": str(r[6]).strip() if r[6] else "-",
                "nd": str(r[8]).strip() if r[8] else "-",
                "pi": str(r[10]).strip() if r[10] else "-",
                "tipo": str(r[19]).strip() if r[19] else "-",
                "objeto": desc,
                "processo": m_proc.group(1).upper() if m_proc else "-",
                "nc_origem": m_nc.group(1) if m_nc else "-",
                "status_assinatura": str(r[22]).strip() if r[22] else "-",
                "valor_empenhado": valor,
                "situacao": "ANULADA" if abs(valor) < 0.005 else "EMPENHADA",
                "valor_liquidado": None,
                "saldo_a_liquidar": None,
                "itens": [],
            }
            cur = key
        elif cur:
            desc_item = str(r[13]).strip() if r[13] else ""
            if desc_item and len(nes[cur]["itens"]) < 8:
                nes[cur]["itens"].append({
                    "item_num": str(int(num(r[12]))) if r[12] is not None else "-",
                    "descricao": desc_item,
                    "operacao": str(r[17]).strip() if r[17] else "-",
                    "valor": round(num(r[16]), 2),
                })
    wb.close()
    return list(nes.values())


def extrair_rpnp():
    """Colunas reais: 4=531110100 inscrito | 5=531210000 reinscrição | 6=631100000 a liquidar |
    7=631200000 em liquidação | 8=631300000 liquidado a pagar | 9=631400000 pago."""
    wb = openpyxl.load_workbook(RPNP_XLSX, data_only=True, read_only=True)
    ws = wb.active
    out, ug = [], None
    for r in ws.iter_rows(values_only=True):
        if r[0] and str(r[0]).strip().isdigit():
            ug = str(r[0]).strip()
        if ug not in UGS or not r[2] or str(r[0]).strip() == "Total":
            continue
        if str(r[0]).strip().lower() == "total":
            continue
        out.append({
            "ug": ug, "ug_nome": UG_NOME[ug], "acao": str(r[2]).strip(), "nome": str(r[3]).strip().title(),
            "inscrito": round(num(r[4]), 2), "reinscricao": round(num(r[5]), 2),
            "a_liquidar": round(num(r[6]), 2), "em_liquidacao": round(num(r[7]), 2),
            "liquidado_a_pagar": round(num(r[8]), 2), "pago": round(num(r[9]), 2),
        })
    wb.close()
    return out


def carregar_totais_ug():
    """Totais oficiais SIAFI por UG (conferidos: empenhado fecha com a soma da planilha)."""
    d = json.loads(TOTAIS_JSON.read_text(encoding="utf-8"))
    return d


def gerar_base_completa():
    planilha, origem = obter_planilha_empenhos()
    corte = datetime.date.today()
    empenhos = extrair_empenhos(planilha, corte)
    rap = extrair_rpnp()
    totais_ug = carregar_totais_ug()

    # Conferência: soma das NEs da planilha x empenhado oficial do SIAFI
    conf = {}
    for ug in UGS:
        soma = round(sum(e["valor_empenhado"] for e in empenhos if e["ug"] == ug), 2)
        oficial = totais_ug[ug]["empenhado"]
        conf[ug] = {"soma_planilha": soma, "oficial_siafi": oficial, "diferenca": round(soma - oficial, 2)}

    cons = {k: round(sum(totais_ug[u][k] for u in UGS), 2)
            for k in ("provisao", "credito_disponivel", "empenhado", "liquidado", "pago")}
    cons["a_liquidar"] = round(cons["empenhado"] - cons["liquidado"], 2)
    cons["a_pagar"] = round(cons["liquidado"] - cons["pago"], 2)
    cons["taxa_liquidacao"] = round(cons["liquidado"] / cons["empenhado"] * 100, 2) if cons["empenhado"] else 0.0
    totais = {}
    for ug in UGS:
        t = dict(totais_ug[ug])
        t["a_liquidar"] = round(t["empenhado"] - t["liquidado"], 2)
        t["a_pagar"] = round(t["liquidado"] - t["pago"], 2)
        t["taxa_liquidacao"] = round(t["liquidado"] / t["empenhado"] * 100, 2) if t["empenhado"] else 0.0
        totais[ug] = t
    totais["consolidado"] = cons

    totais_rap = {
        "rpnp_inscrito": round(sum(r["inscrito"] for r in rap), 2),
        "rpnp_a_liquidar": round(sum(r["a_liquidar"] for r in rap), 2),
        "rpnp_em_liquidacao": round(sum(r["em_liquidacao"] for r in rap), 2),
        "rpnp_liquidado_a_pagar": round(sum(r["liquidado_a_pagar"] for r in rap), 2),
        "rpnp_pago": round(sum(r["pago"] for r in rap), 2),
    }

    # Idade dos empenhos emitidos (NÃO é aging de liquidação: a fonte não baixa por liquidação)
    faixas = {"0-30 dias": [0, 30], "31-60 dias": [31, 60], "61-90 dias": [61, 90], "> 90 dias": [91, 10**6]}
    idade = {k: {"qtd": 0, "valor": 0.0} for k in faixas}
    for e in empenhos:
        if e["situacao"] != "EMPENHADA":
            continue
        for k, (a, b) in faixas.items():
            if a <= e["idade_dias"] <= b:
                e["faixa_idade"] = k
                idade[k]["qtd"] += 1
                idade[k]["valor"] = round(idade[k]["valor"] + e["valor_empenhado"], 2)
                break

    forn = {}
    for e in empenhos:
        if e["situacao"] != "EMPENHADA":
            continue
        f = forn.setdefault(e["cnpj"], {"cnpj": e["cnpj"], "razao": e["fornecedor"], "qtd_nes": 0, "tot_empenhado": 0.0})
        f["qtd_nes"] += 1
        f["tot_empenhado"] = round(f["tot_empenhado"] + e["valor_empenhado"], 2)
    top = sorted(forn.values(), key=lambda x: -x["tot_empenhado"])[:15]

    qtd_anuladas = sum(1 for e in empenhos if e["situacao"] == "ANULADA")
    payload = {
        "metadata": {
            "unidade": "Batalhão Central de Manutenção e Suprimento (BCMS)",
            "orgao_vinculado": "Base de Apoio Logístico (Ba Ap Log)",
            "exercicio": 2026,
            "corte_atualizacao": corte.strftime("%d/%m/%Y"),
            "fonte_empenhos": origem,
            "fonte_rpnp": f"Cache do Tesouro Gerencial de {RPNP_DATA_REF}",
            "gerado_em": datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        },
        "auditoria": {
            "conferencia_empenhado": conf,
            "qtd_nes": len(empenhos),
            "qtd_nes_anuladas_valor_zero": qtd_anuladas,
            "liquidacao_por_ne_disponivel": False,
            "limitacao": "A planilha EMPENHO traz as contas 622110000 e 622130100, que não discriminam a liquidação por NE. "
                         "Para identificar NEs já liquidadas é preciso incluir as contas 622920101/622920103 na consulta do Tesouro Gerencial.",
        },
        "totais_siafi": totais,
        "totais_rap": totais_rap,
        "idade_distribuicao": idade,
        "top_fornecedores": top,
        "restos_a_pagar": rap,
        "empenhos": empenhos,
    }
    OUTPUT_JSON.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"{len(empenhos)} NEs ({qtd_anuladas} anuladas) | conferência empenhado: {conf}")
    print(f"RPNP a liquidar: R$ {totais_rap['rpnp_a_liquidar']:.2f}")
    return payload


if __name__ == "__main__":
    gerar_base_completa()
